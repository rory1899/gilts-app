#!/usr/bin/env python3
"""
UK Gilt & Fixed Notice Account - Annual Cash Flow Calculator (with income tax)
------------------------------------------------------------------------------
Pure standard-library (tkinter) - no extra packages required.

GILT ASSUMPTIONS
  * Inputs: ask (clean) price per £100 nominal, coupon %, maturity date and the
    total £ amount you pay.
  * Nominal held  = amount purchased / (ask price / 100)
  * Redemption    = nominal (gilts redeem at par, £100 per £100 nominal)
  * Annual coupon = nominal x coupon %, paid half-yearly. Coupon dates are
    worked back from the maturity date in 6-month steps; the final coupon is
    paid on the maturity date together with the redemption.
  * Accrued interest paid at purchase is ignored (only future cash flows are
    counted).

FIXED NOTICE ACCOUNT ASSUMPTIONS
  * Capital and simple gross interest (balance x rate x days / 365) are paid as
    ONE single payment on the final day of the notice period.
  * Optionally, capital can be rolled into a new term at the same rate
    (interest is then paid out at the end of each term).

FEE ASSUMPTIONS (separate outgoings, deducted from total returns)
  * Purchase fee, charged once per bond on the amount purchased (on top of it),
    assuming purchase on the "cash flows from" date:
      0.10% if the bond has LESS than one year to maturity,
      0.25% if it has one year or more to maturity.
  * Platform fee: 0.15% p.a. of total bond value, paid in arrears at the end of
    each calendar quarter (31 Mar, 30 Jun, 30 Sep, 31 Dec). Bond value is taken
    as the amount purchased (no market re-pricing). Each fee is pro-rated by
    days held in the quarter (first quarter from purchase, last quarter to
    maturity). Fees apply to gilts only, not notice accounts.
  * Fees are not deducted for income tax purposes.
  * Edit the FEES dictionary below if the charges change.

TAX ASSUMPTIONS (rest of UK - savings income is taxed at the same rates in
Scotland, Wales and England)
  * The holding is joint, owned in the proportions you set (default 50/50).
  * Each person has NO other income, so each gets: Personal Allowance, the
    £5,000 starting rate for savings, and the £1,000 Personal Savings Allowance.
  * Gilt coupons and notice-account interest are both "savings income" and are
    taxed in the UK tax year (6 April - 5 April) in which they are paid.
  * Redemption of gilts and return of account capital are NOT taxed. Gilts are
    exempt from Capital Gains Tax, so a gain from buying below par is tax free.
  * Savings tax rates are 20/40/45% up to 5 April 2027 and 22/42/47% from
    6 April 2027 (Finance Act 2026). Allowances are assumed frozen.
  * This is a planning estimate, not tax advice.
"""

import calendar
import csv
import tkinter as tk
from datetime import date, datetime
from tkinter import filedialog, messagebox, ttk

GBP = lambda x: f"£{x:,.2f}"

# --------------------------------------------------------------------------
# UK tax rules (edit here if allowances / rates change)
# --------------------------------------------------------------------------
TAX_RULES = {
    "personal_allowance": 12_570,
    "pa_taper_start": 100_000,        # PA reduced £1 for every £2 above this
    "starting_rate_band": 5_000,      # 0% starting rate for savings
    "basic_band": 37_700,             # taxable income taxed at basic rate
    "additional_threshold": 125_140,  # taxable income above this = additional rate
    "psa_basic": 1_000,               # Personal Savings Allowance
    "psa_higher": 500,
    "psa_additional": 0,
}


FEES = {
    "short_rate": 0.0010,      # purchase fee, < 1 year to maturity
    "long_rate": 0.0025,       # purchase fee, >= 1 year to maturity
    "platform_rate": 0.0015,   # platform fee p.a. on total bond value
}
FEE_TYPES = ("purchase_fee", "platform_fee")


def savings_rates(tax_year_start: int):
    """(basic, higher, additional) savings tax rates for a tax year."""
    if tax_year_start >= 2027:
        return 0.22, 0.42, 0.47
    return 0.20, 0.40, 0.45


def tax_year_start(d: date) -> int:
    return d.year if (d.month, d.day) >= (4, 6) else d.year - 1


def tax_year_label(y: int) -> str:
    return f"{y}/{(y + 1) % 100:02d}"


def savings_tax(income: float, ty: int) -> dict:
    """Income tax for ONE person whose only income is `income` of savings
    interest in tax year starting `ty`."""
    r = TAX_RULES
    pa = r["personal_allowance"]
    if income > r["pa_taper_start"]:
        pa = max(0.0, pa - (income - r["pa_taper_start"]) / 2)
    pa_used = min(income, pa)
    taxable = income - pa_used  # taxable income after personal allowance

    srs_used = min(taxable, r["starting_rate_band"])

    if taxable > r["additional_threshold"]:
        psa = r["psa_additional"]
    elif taxable > r["basic_band"]:
        psa = r["psa_higher"]
    else:
        psa = r["psa_basic"]
    psa_used = min(taxable - srs_used, psa)

    nil_end = srs_used + psa_used  # position in taxable income where 0% ends
    b, h, a = savings_rates(ty)
    basic_amt = max(0.0, min(taxable, r["basic_band"]) - nil_end)
    higher_amt = max(0.0, min(taxable, r["additional_threshold"])
                     - max(nil_end, r["basic_band"]))
    add_amt = max(0.0, taxable - max(nil_end, r["additional_threshold"]))
    tax = basic_amt * b + higher_amt * h + add_amt * a
    return {"income": income, "pa_used": pa_used, "srs_used": srs_used,
            "psa_used": psa_used, "taxed_amount": taxable - nil_end, "tax": tax}


# --------------------------------------------------------------------------
# Helpers
# --------------------------------------------------------------------------
def parse_date(text: str) -> date:
    text = text.strip()
    for fmt in ("%d/%m/%Y", "%d/%m/%y", "%Y-%m-%d", "%d-%m-%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"Invalid date '{text}'. Use DD/MM/YYYY.")


def parse_num(text: str, field: str) -> float:
    try:
        return float(text.replace("£", "").replace(",", "").replace("%", "").strip())
    except ValueError:
        raise ValueError(f"'{field}' must be a number.")


def add_months(d: date, n: int) -> date:
    y, m = divmod(d.month - 1 + n, 12)
    year, month = d.year + y, m + 1
    return date(year, month, min(d.day, calendar.monthrange(year, month)[1]))


def fmt_date(d: date) -> str:
    return d.strftime("%d/%m/%Y")


# --------------------------------------------------------------------------
# Cash flow engine
# --------------------------------------------------------------------------
def gilt_events(g, start, end):
    ev = []
    mat = g["maturity"]
    half = g["annual_coupon"] / 2
    k = 0
    while True:
        cd = add_months(mat, -6 * k)
        if cd <= start:
            break
        if cd <= end:
            ev.append((cd, "coupon", half, f"{g['name']} coupon"))
        k += 1
    if start < mat <= end:
        ev.append((mat, "redemption", g["nominal"], f"{g['name']} redemption"))
    return ev


def account_events(a, start, end):
    ev = []
    bal = a["amount"]
    cur = a["start"]
    for _ in range(1000):
        if a["unit"] == "months":
            term_end = add_months(cur, a["length"])
        else:
            term_end = date.fromordinal(cur.toordinal() + a["length"])
        if term_end > end:
            break
        interest = bal * a["rate"] / 100 * (term_end - cur).days / 365
        if term_end > start:
            ev.append((term_end, "interest", interest, f"{a['name']} interest"))
        if a["rollover"]:
            cur = term_end
        else:
            if term_end > start:
                ev.append((term_end, "capital", bal, f"{a['name']} capital returned"))
            break
    return ev


def purchase_fee_rate(g, start):
    """0.10% if less than one year to maturity from purchase, else 0.25%."""
    return FEES["short_rate"] if g["maturity"] < add_months(start, 12) else FEES["long_rate"]


def purchase_fee_events(gilts, start):
    ev = []
    for g in gilts:
        if g["maturity"] <= start:
            continue
        rate = purchase_fee_rate(g, start)
        ev.append((start, "purchase_fee", g["amount"] * rate,
                   f"{g['name']} purchase fee ({rate:.2%})"))
    return ev


def quarter_ends(start, end):
    ends = []
    for y in range(start.year - 1, end.year + 2):
        for m in (3, 6, 9, 12):
            ends.append(date(y, m, calendar.monthrange(y, m)[1]))
    return ends


def platform_fee_events(gilts, start, end):
    """One fee per quarter end, summed across gilts held in that quarter."""
    ev = []
    qe = quarter_ends(start, end)
    for prev, q in zip(qe, qe[1:]):
        if q <= start or q > end:
            continue
        fee = 0.0
        for g in gilts:
            if g["maturity"] <= start:
                continue
            days = (min(q, g["maturity"]) - max(prev, start)).days
            if days > 0:
                fee += g["amount"] * FEES["platform_rate"] * days / 365
        if fee > 0:
            ev.append((q, "platform_fee", fee,
                       f"Platform fee, quarter ending {fmt_date(q)}"))
    return ev


SUMMARY_HEADS = ("Tax year", "Gilt coupons", "Account interest", "Capital returned",
                 "Gross income", "Income tax", "Net income (after tax)",
                 "Purchase fees", "Platform fees", "Total net cash")


def build_returns(gilts, accounts, start, end, coupons, interest, tax,
                  purchase_fees, platform_fees):
    """Total-return breakdown. Returns list of (label, amount|None, pct|None, bold)
    and a dict of key figures. Gilts maturing after the period are held at cost
    (their capital gain/loss is not yet realised)."""
    live = [g for g in gilts if g["maturity"] > start]
    gilt_cost = sum(g["amount"] for g in live)
    acc_cost = sum(a["amount"] for a in accounts)
    capital = gilt_cost + acc_cost
    redeemed = [g for g in live if g["maturity"] <= end]
    gain_lines = [(g, g["nominal"] - g["amount"]) for g in redeemed]
    gain = sum(x for _, x in gain_lines)
    gross = interest + coupons + gain
    fees = purchase_fees + platform_fees
    net = gross - tax - fees
    pct = lambda v: (v / capital) if capital else None

    rows = [
        (f"TOTAL RETURNS OVER THE PERIOD ({fmt_date(start)} to {fmt_date(end)})", None, None, True),
        ("INITIAL CAPITAL INVESTED", capital, None, True),
        ("   Gilts (total amount purchased)", gilt_cost, None, False),
        ("   Fixed notice accounts (deposits)", acc_cost, None, False),
        ("RETURN ON TOP OF CAPITAL", None, None, True),
        ("Interest gained (fixed notice accounts)", interest, pct(interest), False),
        ("Coupon payments (gilts)", coupons, pct(coupons), False),
        ("Capital gain / (loss) on bond redemptions", gain, pct(gain), False),
    ]
    for g, x in gain_lines:
        rows.append((f"   {g['name']}: redeems {GBP(g['nominal'])} less cost {GBP(g['amount'])}",
                     x, None, False))
    rows += [
        ("TOTAL GROSS RETURN", gross, pct(gross), True),
        ("Less: income tax on interest and coupons", -tax, pct(-tax), False),
        ("Less: bond purchase fees (outgoing)", -purchase_fees, pct(-purchase_fees), False),
        ("Less: quarterly platform fees (outgoing)", -platform_fees, pct(-platform_fees), False),
        ("TOTAL NET RETURN (after income tax and fees)", net, pct(net), True),
        ("INITIAL CAPITAL + NET RETURN", capital + net, None, True),
        ("Note: gilts maturing after the period are held at cost (gain/loss not yet realised); "
         "redemption gains are free of income tax and CGT.", None, None, False),
    ]
    return rows, {"capital": capital, "gain": gain, "gross": gross, "fees": fees, "net": net}


# --------------------------------------------------------------------------
# UI
# --------------------------------------------------------------------------
class AccountFrame(ttk.LabelFrame):
    def __init__(self, master, title, default_name):
        super().__init__(master, text=title, padding=8)
        self.enabled = tk.BooleanVar(value=False)
        self.name = tk.StringVar(value=default_name)
        self.amount = tk.StringVar()
        self.rate = tk.StringVar()
        self.length = tk.StringVar()
        self.unit = tk.StringVar(value="months")
        self.start = tk.StringVar(value=fmt_date(date.today()))
        self.rollover = tk.BooleanVar(value=False)

        ttk.Checkbutton(self, text="Use this account", variable=self.enabled).grid(
            row=0, column=0, columnspan=4, sticky="w")
        rows = [("Name", self.name), ("Amount deposited (£)", self.amount),
                ("Fixed rate (% p.a.)", self.rate), ("Start date", self.start)]
        for i, (lbl, var) in enumerate(rows, start=1):
            ttk.Label(self, text=lbl).grid(row=i, column=0, sticky="w", pady=2)
            ttk.Entry(self, textvariable=var, width=18).grid(
                row=i, column=1, columnspan=3, sticky="ew", pady=2)
        ttk.Label(self, text="Notice length").grid(row=5, column=0, sticky="w", pady=2)
        ttk.Entry(self, textvariable=self.length, width=8).grid(row=5, column=1, sticky="w")
        ttk.Combobox(self, textvariable=self.unit, values=("days", "months"),
                     width=8, state="readonly").grid(row=5, column=2, sticky="w", padx=4)
        ttk.Checkbutton(self, text="Roll capital into new term at same rate\n"
                                   "(interest paid out each term)",
                        variable=self.rollover).grid(row=6, column=0, columnspan=4,
                                                     sticky="w", pady=(4, 0))
        self.columnconfigure(1, weight=1)

    def get(self):
        if not self.enabled.get():
            return None
        nm = self.name.get().strip() or "Account"
        try:
            length = int(self.length.get().strip())
        except ValueError:
            raise ValueError(f"{nm}: notice length must be a whole number.")
        if length <= 0:
            raise ValueError(f"{nm}: notice length must be greater than zero.")
        return {
            "name": nm,
            "amount": parse_num(self.amount.get(), f"{nm} amount"),
            "rate": parse_num(self.rate.get(), f"{nm} rate"),
            "length": length,
            "unit": self.unit.get(),
            "start": parse_date(self.start.get()),
            "rollover": self.rollover.get(),
        }


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("UK Gilt & Fixed Notice Account - Annual Cash Flow")
        self.geometry("1100x860")
        self.minsize(960, 720)
        self.gilts = []
        self.events = []
        self.summary_rows = []
        self.tax_rows = []
        self.return_rows = []
        self._build()

    # ---- layout ----------------------------------------------------------
    def _build(self):
        pad = {"padx": 10, "pady": 5}

        # Gilts ------------------------------------------------------------
        gf = ttk.LabelFrame(self, text="Gilt purchases", padding=8)
        gf.pack(fill="x", **pad)

        self.g_name = tk.StringVar()
        self.g_ask = tk.StringVar()
        self.g_coupon = tk.StringVar()
        self.g_mat = tk.StringVar()
        self.g_amount = tk.StringVar()

        fields = [("Gilt name", self.g_name, 18),
                  ("Ask price (per £100)", self.g_ask, 12),
                  ("Coupon (%)", self.g_coupon, 10),
                  ("Maturity (DD/MM/YYYY)", self.g_mat, 14),
                  ("Total amount purchased (£)", self.g_amount, 16)]
        for i, (lbl, var, w) in enumerate(fields):
            ttk.Label(gf, text=lbl).grid(row=0, column=i, sticky="w", padx=3)
            ttk.Entry(gf, textvariable=var, width=w).grid(row=1, column=i, padx=3, sticky="w")
        ttk.Button(gf, text="Add gilt", command=self.add_gilt).grid(row=1, column=5, padx=8)
        ttk.Button(gf, text="Remove selected", command=self.remove_gilt).grid(row=1, column=6)

        cols = ("name", "ask", "coupon", "maturity", "amount", "nominal", "annual")
        self.gilt_tree = ttk.Treeview(gf, columns=cols, show="headings", height=4)
        heads = ("Gilt", "Ask price", "Coupon", "Maturity", "Amount purchased",
                 "Nominal = redemption", "Annual coupon (£)")
        widths = (170, 80, 70, 100, 130, 140, 130)
        for c, h, w in zip(cols, heads, widths):
            self.gilt_tree.heading(c, text=h)
            self.gilt_tree.column(c, width=w, anchor="w" if c == "name" else "e")
        self.gilt_tree.grid(row=2, column=0, columnspan=7, sticky="ew", pady=(8, 0))
        ttk.Label(gf, text="Nominal and redemption are inferred from the ask price: "
                           "amount ÷ (ask ÷ 100). Redemption is at par (£100).",
                  foreground="#555").grid(row=3, column=0, columnspan=7, sticky="w", pady=(4, 0))
        gf.columnconfigure(6, weight=1)

        # Accounts ---------------------------------------------------------
        af = ttk.Frame(self)
        af.pack(fill="x", **pad)
        self.acc1 = AccountFrame(af, "Fixed notice account 1", "Notice account 1")
        self.acc2 = AccountFrame(af, "Fixed notice account 2", "Notice account 2")
        self.acc1.pack(side="left", fill="both", expand=True, padx=(0, 5))
        self.acc2.pack(side="left", fill="both", expand=True, padx=(5, 0))

        # Options ----------------------------------------------------------
        of = ttk.LabelFrame(self, text="Projection & tax options", padding=8)
        of.pack(fill="x", **pad)
        self.from_date = tk.StringVar(value=fmt_date(date.today()))
        self.years = tk.StringVar(value="5")
        self.use_tax = tk.BooleanVar(value=True)
        self.split = tk.StringVar(value="50")
        ttk.Label(of, text="Cash flows from").grid(row=0, column=0, sticky="w")
        ttk.Entry(of, textvariable=self.from_date, width=12).grid(row=0, column=1, padx=5)
        ttk.Label(of, text="Years to project").grid(row=0, column=2, padx=(15, 0))
        ttk.Entry(of, textvariable=self.years, width=5).grid(row=0, column=3, padx=5)
        ttk.Checkbutton(of, text="Deduct UK income tax (joint holding)",
                        variable=self.use_tax).grid(row=0, column=4, padx=(20, 5))
        ttk.Label(of, text="Person 1 owns (%)").grid(row=0, column=5, padx=(10, 0))
        ttk.Entry(of, textvariable=self.split, width=5).grid(row=0, column=6, padx=5)
        ttk.Label(of, text="(Person 2 owns the rest; each has no other income)",
                  foreground="#555").grid(row=1, column=4, columnspan=4, sticky="w", padx=(20, 0))

        bf = ttk.Frame(self)
        bf.pack(fill="x", padx=10)
        ttk.Button(bf, text="Calculate cash flow", command=self.calculate).pack(side="left")
        ttk.Button(bf, text="Export CSV", command=self.export_csv).pack(side="left", padx=8)
        ttk.Button(bf, text="Export Excel (bold headings)",
                   command=self.export_xlsx).pack(side="left")
        ttk.Button(bf, text="Clear results", command=self.clear_results).pack(side="left", padx=8)

        # Results ----------------------------------------------------------
        nb = ttk.Notebook(self)
        nb.pack(fill="both", expand=True, **pad)

        t0 = ttk.Frame(nb)
        nb.add(t0, text="Total returns")
        self.ret_tree = ttk.Treeview(t0, columns=("item", "amount", "pct"), show="headings")
        for c, h, w, an in (("item", "Item", 560, "w"), ("amount", "Amount", 140, "e"),
                            ("pct", "% of initial capital", 150, "e")):
            self.ret_tree.heading(c, text=h)
            self.ret_tree.column(c, width=w, anchor=an)
        self.ret_tree.tag_configure("bold", font=("TkDefaultFont", 10, "bold"))
        self.ret_tree.pack(fill="both", expand=True)

        t1 = ttk.Frame(nb)
        nb.add(t1, text="Annual summary (UK tax years)")
        scols = ("period", "coupons", "interest", "capital", "gross", "tax", "net",
                 "pfee", "plat", "total")
        self.sum_tree = ttk.Treeview(t1, columns=scols, show="headings")
        heads = tuple(h + "*" if h == "Capital returned" else h for h in SUMMARY_HEADS)
        widths = (70, 100, 105, 110, 100, 90, 125, 95, 95, 115)
        for c, h, w in zip(scols, heads, widths):
            self.sum_tree.heading(c, text=h)
            self.sum_tree.column(c, width=w, anchor="w" if c == "period" else "e")
        self.sum_tree.tag_configure("total", font=("TkDefaultFont", 10, "bold"))
        self.sum_tree.pack(fill="both", expand=True)
        ttk.Label(t1, text="*Gilt redemptions + notice account capital returned (not taxable). "
                           "Gross income = coupons + interest. "
                           "Fees are separate outgoings, subtracted in 'Total net cash'.",
                  foreground="#555").pack(anchor="w")

        t2 = ttk.Frame(nb)
        nb.add(t2, text="Tax breakdown")
        tcols = ("ty", "person", "income", "pa", "srs", "psa", "taxed", "tax")
        self.tax_tree = ttk.Treeview(t2, columns=tcols, show="headings")
        heads = ("Tax year", "Person", "Savings income", "Personal allowance used",
                 "0% starting rate used", "Savings allowance used",
                 "Taxed at savings rates", "Income tax")
        widths = (80, 70, 115, 150, 140, 140, 140, 100)
        for c, h, w in zip(tcols, heads, widths):
            self.tax_tree.heading(c, text=h)
            self.tax_tree.column(c, width=w, anchor="w" if c in ("ty", "person") else "e")
        self.tax_tree.pack(fill="both", expand=True)

        t3 = ttk.Frame(nb)
        nb.add(t3, text="Payment schedule")
        pcols = ("date", "desc", "type", "amount")
        self.pay_tree = ttk.Treeview(t3, columns=pcols, show="headings")
        for c, h, w in zip(pcols, ("Date", "Description", "Type", "Amount"),
                           (110, 300, 120, 120)):
            self.pay_tree.heading(c, text=h)
            self.pay_tree.column(c, width=w, anchor="e" if c == "amount" else "w")
        self.pay_tree.pack(fill="both", expand=True)

        self.status = tk.StringVar(
            value="Add gilts and/or enable a notice account, then press Calculate.")
        ttk.Label(self, textvariable=self.status, foreground="#555").pack(
            anchor="w", padx=10, pady=(0, 6))

    # ---- gilt list -------------------------------------------------------
    def add_gilt(self):
        try:
            ask = parse_num(self.g_ask.get(), "Ask price")
            coupon = parse_num(self.g_coupon.get(), "Coupon %")
            amount = parse_num(self.g_amount.get(), "Total amount purchased")
            mat = parse_date(self.g_mat.get())
            if ask <= 0:
                raise ValueError("Ask price must be greater than zero.")
            if amount <= 0:
                raise ValueError("Total amount purchased must be greater than zero.")
            if coupon < 0:
                raise ValueError("Coupon cannot be negative.")
        except ValueError as e:
            messagebox.showerror("Invalid input", str(e))
            return
        nominal = amount / (ask / 100)
        name = self.g_name.get().strip() or f"Gilt {len(self.gilts) + 1}"
        self.gilts.append({"name": name, "ask": ask, "coupon_pct": coupon,
                           "maturity": mat, "amount": amount, "nominal": nominal,
                           "annual_coupon": nominal * coupon / 100})
        for v in (self.g_name, self.g_ask, self.g_coupon, self.g_mat, self.g_amount):
            v.set("")
        self._refresh_gilts()

    def remove_gilt(self):
        for idx in sorted((self.gilt_tree.index(i) for i in self.gilt_tree.selection()),
                          reverse=True):
            del self.gilts[idx]
        self._refresh_gilts()

    def _refresh_gilts(self):
        self.gilt_tree.delete(*self.gilt_tree.get_children())
        for g in self.gilts:
            self.gilt_tree.insert("", "end", values=(
                g["name"], f"{g['ask']:.3f}", f"{g['coupon_pct']:.3f}%",
                fmt_date(g["maturity"]), GBP(g["amount"]),
                GBP(g["nominal"]), GBP(g["annual_coupon"])))

    # ---- calculation -----------------------------------------------------
    def calculate(self):
        try:
            start = parse_date(self.from_date.get())
            years = int(self.years.get().strip())
            if years <= 0:
                raise ValueError("Years to project must be positive.")
            share1 = parse_num(self.split.get(), "Person 1 ownership %")
            if not 0 <= share1 <= 100:
                raise ValueError("Person 1 ownership must be between 0 and 100.")
            accounts = [a for a in (self.acc1.get(), self.acc2.get()) if a]
        except ValueError as e:
            messagebox.showerror("Invalid input", str(e))
            return
        if not self.gilts and not accounts:
            messagebox.showinfo("Nothing to calculate",
                                "Add at least one gilt or enable a notice account.")
            return

        end = add_months(start, 12 * years)
        events = []
        for g in self.gilts:
            events += gilt_events(g, start, end)
        for a in accounts:
            events += account_events(a, start, end)
        events += purchase_fee_events(self.gilts, start)
        events += platform_fee_events(self.gilts, start, end)
        events.sort(key=lambda e: e[0])
        self.events = events

        # aggregate by UK tax year
        agg = {}
        for d, typ, amt, _ in events:
            y = tax_year_start(d)
            row = agg.setdefault(y, {"coupon": 0.0, "interest": 0.0,
                                     "redemption": 0.0, "capital": 0.0,
                                     "purchase_fee": 0.0, "platform_fee": 0.0})
            row[typ] += amt

        use_tax = self.use_tax.get()
        shares = (share1 / 100, 1 - share1 / 100)

        self.sum_tree.delete(*self.sum_tree.get_children())
        self.tax_tree.delete(*self.tax_tree.get_children())
        self.summary_rows, self.tax_rows = [], []
        T = dict.fromkeys(("coupon", "interest", "capital", "gross", "tax", "net",
                           "pfee", "plat", "total"), 0.0)

        for y in sorted(agg):
            r = agg[y]
            gross = r["coupon"] + r["interest"]
            capital = r["redemption"] + r["capital"]
            tax = 0.0
            if use_tax:
                for p, sh in enumerate(shares, start=1):
                    res = savings_tax(gross * sh, y)
                    tax += res["tax"]
                    trow = (tax_year_label(y), f"Person {p}", GBP(res["income"]),
                            GBP(res["pa_used"]), GBP(res["srs_used"]), GBP(res["psa_used"]),
                            GBP(res["taxed_amount"]), GBP(res["tax"]))
                    self.tax_rows.append(trow)
                    self.tax_tree.insert("", "end", values=trow)
            net = gross - tax
            pfee, plat = r["purchase_fee"], r["platform_fee"]
            total = net + capital - pfee - plat
            for k, v in (("coupon", r["coupon"]), ("interest", r["interest"]),
                         ("capital", capital), ("gross", gross), ("tax", tax),
                         ("net", net), ("pfee", pfee), ("plat", plat), ("total", total)):
                T[k] += v
            vals = (tax_year_label(y), GBP(r["coupon"]), GBP(r["interest"]), GBP(capital),
                    GBP(gross), GBP(tax), GBP(net), GBP(-pfee), GBP(-plat), GBP(total))
            self.summary_rows.append(vals)
            self.sum_tree.insert("", "end", values=vals)

        tvals = ("TOTAL", GBP(T["coupon"]), GBP(T["interest"]), GBP(T["capital"]),
                 GBP(T["gross"]), GBP(T["tax"]), GBP(T["net"]), GBP(-T["pfee"]),
                 GBP(-T["plat"]), GBP(T["total"]))
        self.summary_rows.append(tvals)
        self.sum_tree.insert("", "end", values=tvals, tags=("total",))

        self.return_rows, _ = build_returns(self.gilts, accounts, start, end,
                                            T["coupon"], T["interest"], T["tax"],
                                            T["pfee"], T["plat"])
        self.ret_tree.delete(*self.ret_tree.get_children())
        for lbl, amt, pc, bold in self.return_rows:
            self.ret_tree.insert("", "end", tags=("bold",) if bold else (), values=(
                lbl, "" if amt is None else GBP(amt), "" if pc is None else f"{pc:.2%}"))

        labels = {"coupon": "Gilt coupon", "redemption": "Gilt redemption",
                  "interest": "Account interest", "capital": "Account capital",
                  "purchase_fee": "Purchase fee (cost)", "platform_fee": "Platform fee (cost)"}
        self.pay_tree.delete(*self.pay_tree.get_children())
        for d, typ, amt, desc in events:
            shown = -amt if typ in FEE_TYPES else amt
            self.pay_tree.insert("", "end", values=(fmt_date(d), desc, labels[typ], GBP(shown)))

        self.status.set(
            f"{len(events)} cash flows from {fmt_date(start)} to {fmt_date(end)}. "
            + ("Tax estimated per person using UK savings-income rules (planning estimate, "
               "not tax advice). First/last tax years may be partial." if use_tax
               else "Tax not deducted."))

    def clear_results(self):
        for t in (self.ret_tree, self.sum_tree, self.tax_tree, self.pay_tree):
            t.delete(*t.get_children())
        self.events, self.summary_rows, self.tax_rows, self.return_rows = [], [], [], []

    def export_csv(self):
        if not self.summary_rows:
            messagebox.showinfo("Nothing to export", "Run a calculation first.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".csv",
                                            filetypes=[("CSV", "*.csv")])
        if not path:
            return
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f)
            # Total returns block first (CSV cannot store bold, so key rows are
            # in CAPITALS; use "Export Excel" for real bold formatting).
            for lbl, amt, pc, bold in self.return_rows:
                w.writerow((lbl, "" if amt is None else f"{amt:.2f}",
                            "" if pc is None else f"{pc:.2%}"))
            w.writerow([])
            w.writerow(SUMMARY_HEADS)
            w.writerows(self.summary_rows)
            fee_ev = [e for e in self.events if e[1] in FEE_TYPES]
            w.writerow([])
            w.writerow(("FEES (OUTGOING COSTS)",))
            w.writerow(("Date", "Description", "Type", "Cost"))
            for d, typ, amt, desc in fee_ev:
                w.writerow((fmt_date(d), desc, typ.replace("_", " "), f"{amt:.2f}"))
            w.writerow(("TOTAL FEES", "", "", f"{sum(e[2] for e in fee_ev):.2f}"))
            if self.tax_rows:
                w.writerow([])
                w.writerow(("Tax year", "Person", "Savings income", "Personal allowance used",
                            "0% starting rate used", "Savings allowance used",
                            "Taxed at savings rates", "Income tax"))
                w.writerows(self.tax_rows)
            w.writerow([])
            w.writerow(("Date", "Description", "Type", "Amount"))
            for d, typ, amt, desc in self.events:
                signed = -amt if typ in FEE_TYPES else amt
                w.writerow((fmt_date(d), desc, typ, f"{signed:.2f}"))
        self.status.set(f"Exported to {path}")

    def export_xlsx(self):
        if not self.summary_rows:
            messagebox.showinfo("Nothing to export", "Run a calculation first.")
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font
        except ImportError:
            messagebox.showerror("openpyxl needed",
                                 "Excel export needs openpyxl:\n\npip install openpyxl\n\n"
                                 "Alternatively use Export CSV.")
            return
        path = filedialog.asksaveasfilename(defaultextension=".xlsx",
                                            filetypes=[("Excel workbook", "*.xlsx")])
        if not path:
            return
        num = lambda t: float(t.replace("£", "").replace(",", ""))
        bold = Font(bold=True)
        money = '"£"#,##0.00;[Red]-"£"#,##0.00'

        wb = Workbook()
        ws = wb.active
        ws.title = "Total returns"
        ws.append(["Item", "Amount", "% of initial capital"])
        for c in ws[1]:
            c.font = bold
        for lbl, amt, pc, b in self.return_rows:
            ws.append([lbl, amt, pc])
            r = ws.max_row
            ws.cell(r, 2).number_format = money
            ws.cell(r, 3).number_format = "0.00%"
            if b:
                for c in range(1, 4):
                    ws.cell(r, c).font = bold
        ws.column_dimensions["A"].width = 70
        ws.column_dimensions["B"].width = 16
        ws.column_dimensions["C"].width = 20

        ws2 = wb.create_sheet("Annual summary")
        ws2.append(list(SUMMARY_HEADS))
        for row in self.summary_rows:
            ws2.append([row[0]] + [num(x) for x in row[1:]])
            for c in range(2, 11):
                ws2.cell(ws2.max_row, c).number_format = money
        for c in ws2[1] + ws2[ws2.max_row]:
            c.font = bold
        for col in "ABCDEFGHIJ":
            ws2.column_dimensions[col].width = 20

        if self.tax_rows:
            ws3 = wb.create_sheet("Tax breakdown")
            ws3.append(["Tax year", "Person", "Savings income", "Personal allowance used",
                        "0% starting rate used", "Savings allowance used",
                        "Taxed at savings rates", "Income tax"])
            for row in self.tax_rows:
                ws3.append(list(row[:2]) + [num(x) for x in row[2:]])
                for c in range(3, 9):
                    ws3.cell(ws3.max_row, c).number_format = money
            for c in ws3[1]:
                c.font = bold
            for col in "ABCDEFGH":
                ws3.column_dimensions[col].width = 22

        fee_ev = [e for e in self.events if e[1] in FEE_TYPES]
        ws5 = wb.create_sheet("Fees (outgoings)")
        ws5.append(["Date", "Description", "Type", "Cost"])
        for d, typ, amt, desc in fee_ev:
            ws5.append([d, desc, typ.replace("_", " "), amt])
            ws5.cell(ws5.max_row, 1).number_format = "dd/mm/yyyy"
            ws5.cell(ws5.max_row, 4).number_format = money
        ws5.append(["TOTAL FEES", None, None, sum(e[2] for e in fee_ev)])
        ws5.cell(ws5.max_row, 4).number_format = money
        for c in list(ws5[1]) + list(ws5[ws5.max_row]):
            c.font = bold
        ws5.column_dimensions["A"].width = 14
        ws5.column_dimensions["B"].width = 44
        ws5.column_dimensions["C"].width = 16
        ws5.column_dimensions["D"].width = 14

        ws4 = wb.create_sheet("Payment schedule")
        ws4.append(["Date", "Description", "Type", "Amount"])
        for d, typ, amt, desc in self.events:
            ws4.append([d, desc, typ, -amt if typ in FEE_TYPES else amt])
            ws4.cell(ws4.max_row, 1).number_format = "dd/mm/yyyy"
            ws4.cell(ws4.max_row, 4).number_format = money
        for c in ws4[1]:
            c.font = bold
        ws4.column_dimensions["B"].width = 36
        ws4.column_dimensions["A"].width = 12
        ws4.column_dimensions["D"].width = 14

        wb.save(path)
        self.status.set(f"Exported to {path}")


if __name__ == "__main__":
    App().mainloop()
